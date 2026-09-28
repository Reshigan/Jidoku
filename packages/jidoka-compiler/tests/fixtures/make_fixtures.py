"""Build the mobilisation-pack fixtures: the real pack's shape, invented values.

The pack itself is client material and is never committed. What the tests need is its *shape* — a
three-row banner above the header, a task register with a dependency column, gates with named
approvers, a KT tracker and a deliverables register under different headers, decisions with no
declared type, boundary conditions with consequences, and an alignment matrix. All of that is here.
"""
from openpyxl import Workbook

BANNER = ["Fixture — invented values, same shape as a real mobilisation pack"]
NOTE = ["Yellow columns maintained by the PM."]
FOOT = ["ACME · FIXTURE-001 v1.0 · 1 Oct – 1 Dec 2026"]


def sheet(wb, name, header, rows):
    ws = wb.create_sheet(name)
    ws.append(BANNER); ws.append(NOTE); ws.append(FOOT); ws.append([])
    ws.append(header)
    for r in rows:
        ws.append(r)


wb = Workbook(); wb.remove(wb.active)
ws = wb.create_sheet("Summary"); ws.append(BANNER); ws.append(["A view, not a register."])

sheet(wb, "Tasks",
      ["ID", "Week", "Owner", "Task", "Gate/Control", "Due", "Depends on", "Status", "Note"],
      [["T-001", "B1", "PM", "Access check for all consultants", "DP-B01", "1 Oct", "",
        "Not started", ""],
       ["T-002", "B1", "Lead", "Take the baseline extract", "G1", "3 Oct", "T-001",
        "In progress", ""],
       ["T-003", "B2", "Lead", "Load the picklist option set", "G2", "10 Oct", "T-002; T-001",
        "Complete", ""]])

sheet(wb, "Gates", ["Gate", "Name", "Date", "Criteria", "Evidence", "Approver", "Passed"],
      [["G1", "Baseline accepted", "Fri 9 Oct", "Baseline extract taken and accepted",
        "Baseline manifest; profiling report", "Lead + client IT owner", ""],
       ["G2", "Foundation signed", "Fri 23 Oct", "Foundation objects verified in QAL",
        "Verification report", "Client IT owner", ""],
       ["G3", "Hypercare entry", "week 8", "Support model accepted", "Signed RACI",
        "Client IT owner", ""]])

# A second register in the same sheet, under its own header, after a blank row and a caption —
# exactly as the real pack writes its one-way doors below its gates.
ws = wb["Gates"]
ws.append([])
ws.append(["One-way doors"])
ws.append(["#", "Door", "When", "Control", "Approver", "Executed"])
ws.append(["D1", "Foundation freeze", "Fri 16 Oct", "Design authority sign-off", "Lead", ""])
ws.append(["D2", "Production data load", "Sat 28 Nov", "Two-person rule",
           "Load lead + client IT", ""])

sheet(wb, "Decisions", ["DP", "Decision", "Owner", "GONXT position", "Due", "Status"],
      [["DP-B01", "Which legal entities are in scope", "Client IT owner",
        "The four trading entities only", "3 Oct", "Open"],
       ["DP-B02", "Pay frequency for monthly staff", "Client payroll",
        "Monthly, last working day", "10 Oct", "Open"]])

sheet(wb, "Conditions", ["#", "Condition", "By", "Consequence if it fails"],
      [["1", "DEV and QAL instances provisioned", "1 Oct",
        "Week 1 cannot start; day-for-day slip"],
       ["2", "Legacy read-only access granted", "1 Oct",
        "No baseline extract; G1 cannot pass"]])

sheet(wb, "KT Tracker", ["#", "Week", "With", "Competence check", "Evidence", "Status",
                         "Signed by / date"],
      [["K1", "B1", "Lead", "Explains why the baseline extract is taken first",
        "Verbal check + screenshot", "Not started", ""]])

sheet(wb, "Deliverables", ["#", "Deliverable", "Draft due", "Validated by", "Accepted by",
                           "Final due", "Status"],
      [["D1", "Support model and RACI", "9 Oct", "Walkthrough with the IT owner",
        "Client IT owner", "27 Nov", "Not started"],
       ["D2", "Configuration workbook", "25 Oct", "Design authority review", "Lead", "8 Jan",
        "Not started"]])

wb.save("packages/jidoka-compiler/tests/fixtures/mobilisation_plan.xlsx")

wb = Workbook(); wb.remove(wb.active)
sheet(wb, "Data Ownership",
      ["Data domain", "Owner (system of record)", "Consumers", "Alignment note"],
      [["Employment and job information", "EC", "All modules, ECC, Analytics",
        "One writer; everyone else reads"],
       ["Pay components", "EC Payroll", "EC; Analytics", "Codes agreed at G2"]])
sheet(wb, "Interlocks",
      ["ID", "Source", "Target", "What flows / logic", "Failure mode if not held", "Ctrl", "Week"],
      [["I-01", "EC", "ECP", "Employee master replication",
        "Payroll runs against stale master data", "G2", "B2"],
       ["I-02", "Picklists", "EC", "Option set codes",
        "Fields reject values that look valid in the workbook", "G2", "B2"]])
wb.save("packages/jidoka-compiler/tests/fixtures/alignment_matrix.xlsx")
print("written")


# ---------------------------------------------------------------------------------------------
# A design document, written by hand with zipfile. python-docx is not a dependency of this repo
# and a fixture generator is a poor reason to acquire one: a .docx is a zip with an XML document
# inside it, and what the reader needs is a document with tables in it.
import zipfile

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def esc(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def heading(text):
    return (f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            f'<w:r><w:t>{esc(text)}</w:t></w:r></w:p>')


def para(text):
    return f'<w:p><w:r><w:t>{esc(text)}</w:t></w:r></w:p>'


def table(rows):
    # Each cell's text is split across two runs, as Word does whenever formatting changes mid
    # sentence — a reader that takes only the first `w:t` gets half a requirement.
    def cell(text):
        head, _, tail = text.partition(" ")
        runs = f"<w:r><w:t>{esc(head)}</w:t></w:r>" + (
            f'<w:r><w:t xml:space="preserve"> {esc(tail)}</w:t></w:r>' if tail else "")
        return f"<w:tc><w:p>{runs}</w:p></w:tc>"
    body = "".join("<w:tr>" + "".join(cell(c) for c in row) + "</w:tr>" for row in rows)
    return f"<w:tbl>{body}</w:tbl>"


def document(path, parts):
    body = "".join(parts)
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document {W}><w:body>{body}</w:body></w:document>'
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", RELS)
        z.writestr("word/document.xml", xml)


document("packages/jidoka-compiler/tests/fixtures/design_document.docx", [
    heading("What This Document Governs"),
    para("The design authority owns this document. Most of its reasoning is in sentences like "
         "this one, and none of that reaches the platform."),
    # A metadata table no profile matches, so the reader has something to report as unread.
    table([["Document ID", "FIXTURE-SDD-001"], ["Version", "1.0"]]),

    heading("Scope Catalogue"),
    table([["Item", "Scope", "Module", "Stream", "BRS section / detailed specification"],
           ["G01", "Model company activation", "Platform", "S1 Platform", "§5.1"],
           ["G02", "Time Off", "Time", "S2 Time", "§5.2"]]),

    heading("Requirements"),
    table([["ID", "Requirement", "Rationale", "Ctry", "Wave/Wk", "Fit", "Ctrl"],
           ["BRS-EC-001", "Single instance, four country layers", "The baseline extract "
            "establishes what standard content delivered", "All", "W1", "STD", "C01"],
           ["BRS-EC-002", "One picklist source", "Interface mappings break silently otherwise",
            "All", "W2", "CFG", "C01"]]),
    # A second requirements table: a real BRS writes seventeen of them, one per module.
    table([["ID", "Requirement", "Rationale", "Ctry", "Wave/Wk", "Fit", "Ctrl"],
           ["BRS-TIM-001", "Leave accrual by country", "Statutory minima differ", "ZA", "W3",
            "GAP", "C02"]]),

    heading("Control Objectives"),
    table([["ID", "Control objective", "Owner", "Frequency", "Evidence"],
           ["C01", "Every change is attributable to a signed source", "GONXT (execution)",
            "Per load cycle", "Evidence bundle per cycle"]]),

    heading("Data Ownership"),
    table([["Data domain", "Owner", "Consumers", "Alignment note"],
           ["Employment and job information", "EC", "All modules, ECC", "One writer"]]),

    heading("Interlocks"),
    table([["ID", "Source", "Target", "What flows / logic", "Failure mode", "Stream", "Ctrl"],
           ["I-01", "EC", "ECP", "Employee master replication", "Payroll runs against stale "
            "master data", "S2", "C03"]]),

    heading("Ordering Constraints"),
    table([["By end of", "Must be complete", "Because", "What stalls if late"],
           ["Week 1", "Legacy extracts profiled", "I-19 — cleansing rules are designed against "
            "actual data", "The whole migration"]]),

    heading("The Rules the Design Obeys"),
    table([["#", "Principle", "Rationale, and what the alternative costs"],
           ["P1", "Fit to standard; deviations are a bounded pool", "Model-company content "
            "supplies most of it"]]),
    table([["#", "Rule", "What it prevents, and how it is checked"],
           ["A1", "One picklist source", "Interface mappings break silently when an external "
            "code moves. Checked at G2 and nightly"]]),

    heading("Decisions"),
    table([["DP", "Decision required", "GONXT recommended position", "Owner", "Required by"],
           ["DP-C04", "Variable Pay in scope, or next financial year",
            "Defer to FY2027/28", "Steering committee", "31 Aug 2026"]]),
    table([["DP", "Design dependency", "Interim design position"],
           ["DP-B01", "Provisioning access", "No interim position available"]]),

    heading("Boundary Conditions"),
    table([["Condition", "Why the schedule needs it", "Consequence if it fails"],
           ["Delta pool consumption near zero", "The pool exists for the whole programme",
            "Each deviation consumes days from a stream with no float"]]),
])

# A second document restating one decision differently, which is the defect `duplicates` finds.
document("packages/jidoka-compiler/tests/fixtures/licensing_position.docx", [
    heading("Decisions"),
    table([["DP", "Decision required", "GONXT recommended position", "Owner", "Required by"],
           ["DP-C04", "Variable Pay licence purchase", "Do not purchase for Wave 1",
            "Steering committee", "31 Aug 2026"]]),
])
print("documents written")
